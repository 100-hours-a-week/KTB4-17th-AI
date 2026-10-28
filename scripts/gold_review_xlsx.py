"""사람 검수표(엑셀) → gold_drafts.py build 가 읽는 라벨 시트(CSV). 검수자 파일마다 태깅·추출 시트를 하나씩 만든다.

검수표는 기존 골드셋(calibration+regression 의 label_accuracy)을 사람이 다시 다는 엑셀이다. 시트 구성:
기준선택(P1~P8 합의·검수자 ID·합의 기준 버전), 태깅검수 60, 보류검수 3, 추출검수 24, 세부점수(P5=B 전용), 차원사전.
"검수 상태"가 완료인 행만 옮긴다. 보류·작성중 행은 빈 라벨로 두어 build 에서 unlabeled 로 잡힌다.

  uv run python scripts/gold_drafts.py export --kind tagging    # items.jsonl 이 먼저 있어야 한다
  uv run python scripts/gold_drafts.py export --kind build
  uv run python scripts/gold_review_xlsx.py 검수자A.xlsx 검수자B.xlsx
  uv run python scripts/gold_drafts.py build --kind tagging --work-dir data/gold_work/existing \\
      --labels data/gold_work/existing/labels/tagging_reviewer-A.csv data/gold_work/existing/labels/tagging_reviewer-B.csv \\
      --reviewers reviewer-A,reviewer-B

엑셀의 입력확인 수식은 저장한 프로그램에 따라 값이 없을 수 있어 믿지 않고, 같은 검사를 여기서 다시 한다.
기준(P1~P8)·합의 기준 버전이 검수자끼리 다르면 아무것도 쓰지 않는다. P5=B(문항별 최소·최대 점수)는 gold_drafts.py 가
구간(low·mid·high)만 받으므로 아직 지원하지 않는다.
"""

from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from scripts.gold_drafts import (  # noqa: E402
    EXPORT_WORK_DIR,
    LEVELS,
    _dims,
    _rows,
    build_sheet_rows,
    tagging_sheet_rows,
)

CRITERIA = [f"P{i}" for i in range(1, 9)]
SHEETS = {"tagging": "태깅검수", "hold": "보류검수", "build": "추출검수"}
TAG_VALUES = ("primary", "secondary")
SKIP = ("", "생략")


def _text(value) -> str:
    if isinstance(value, bool):  # 엑셀이 true·false 를 논리값으로 바꿔 저장할 때가 있다
        return "true" if value else "false"
    return "" if value is None else str(value).strip()


def _header_key(value) -> str:
    """'거리 두기\\navoidance' → 'avoidance'. 차원 열이 아니면 한글 제목 그대로."""
    return _text(value).split("\n")[-1].strip()


def read_criteria(wb) -> dict:
    ws = wb["기준선택"]
    choices = {}
    for row in ws.iter_rows(values_only=True):
        if _text(row[0]) in CRITERIA:
            choices[_text(row[0])] = _text(row[5]).upper()
    return {"reviewer": _text(ws["B4"].value), "version": _text(ws["E4"].value), "choices": choices}


def criteria_errors(criteria: dict) -> list[str]:
    errors = []
    if not criteria["reviewer"]:
        errors.append("검수자 ID(기준선택 B4)가 비어 있음")
    if not criteria["version"]:
        errors.append("합의 기준 버전(기준선택 E4)이 비어 있음")
    missing = [p for p in CRITERIA if criteria["choices"].get(p) not in ("A", "B")]
    if missing:
        errors.append(f"기준 미선택: {', '.join(missing)}")
    if criteria["choices"].get("P5") == "B":
        errors.append("P5=B(문항별 최소·최대 점수)는 아직 변환하지 않음 — gold_drafts.py 는 low·mid·high 구간만 받는다")
    return errors


def read_review(ws) -> list[dict]:
    """머리글('문항 ID' 가 있는 행) 아래 행들을 {머리글 키: 값} 으로."""
    rows = list(ws.iter_rows(values_only=True))
    head_at = next(i for i, r in enumerate(rows) if len(r) > 1 and _text(r[1]) == "문항 ID")
    keys = [_header_key(c) for c in rows[head_at]]
    return [dict(zip(keys, r, strict=False)) for r in rows[head_at + 1 :] if len(r) > 1 and _text(r[1])]


def _note(row: dict, evidence_key: str = "차원별 근거 원문") -> str:
    memo, evidence = _text(row.get("메모")), _text(row.get(evidence_key))
    return " | ".join(x for x in (memo, f"근거: {evidence}" if evidence else "") if x)


def tagging_values(row: dict, version: str) -> tuple[dict, list[str], list[str]]:
    """완료 행 → 라벨 시트 칸 {primary, secondary, off_topic, note}, 오류, 경고."""
    scored, textual = _dims()
    errors, warnings = [], []
    if _text(row.get("검수 기준 버전")) != version:
        errors.append(f"검수 기준 버전 {_text(row.get('검수 기준 버전'))!r} ≠ 합의 버전 {version!r}")
    off = _text(row.get("off_topic")).lower()
    if off not in ("true", "false"):
        errors.append("off_topic 미선택")
    picked: dict[str, list[str]] = {"primary": [], "secondary": []}
    for d in scored + textual:
        value = _text(row.get(d)).lower()
        if value in TAG_VALUES:
            picked[value].append(d)
        elif value not in SKIP:
            errors.append(f"{d}={_text(row.get(d))!r} (primary·secondary·생략)")
    labeled = picked["primary"] + picked["secondary"]
    if off == "true" and labeled:
        errors.append("off_topic=true 인데 차원이 선택됨")
    evidence = _text(row.get("차원별 근거 원문"))
    if labeled and not evidence:
        errors.append("근거 원문 없음")
    warnings += [f"근거에 {d} 없음" for d in labeled if evidence and d not in evidence]
    values = {
        "primary": ",".join(picked["primary"]),
        "secondary": ",".join(picked["secondary"]),
        "off_topic": off,
        "note": _note(row),
    }
    return values, errors, warnings


def build_values(row: dict, version: str) -> tuple[dict, list[str], list[str]]:
    """완료 행 → 라벨 시트 칸 {차원: low|mid|high|항목, done, note}, 오류, 경고."""
    scored, textual = _dims()
    errors, warnings = [], []
    if _text(row.get("검수 기준 버전")) != version:
        errors.append(f"검수 기준 버전 {_text(row.get('검수 기준 버전'))!r} ≠ 합의 버전 {version!r}")
    allowed = set(re.split(r"[,\s]+", _text(row.get("입력의 허용 차원")))) - {""}
    values: dict[str, str] = {}
    for d in scored:
        value = _text(row.get(d)).lower()
        if value in SKIP:
            continue
        if d not in allowed:
            errors.append(f"허용 외 차원 {d} 에 값이 있음")
        elif value not in LEVELS:
            errors.append(f"{d}={_text(row.get(d))!r} (low·mid·high·생략)")
        else:
            values[d] = value
    for d in textual:
        value = _text(row.get(d))
        if value in SKIP:
            continue
        if d not in allowed:
            errors.append(f"허용 외 차원 {d} 에 값이 있음")
        else:
            # 줄바꿈으로 나눈 항목은 시트 구분자(;)로. 쉼표는 조건·예외 문구 안에 있을 수 있어 그대로 둔다
            values[d] = ";".join(x.strip() for x in value.splitlines() if x.strip())
    evidence = _text(row.get("차원별 근거 원문"))
    if values and not evidence:
        errors.append("근거 원문 없음")
    warnings += [f"근거에 {d} 없음" for d in values if evidence and d not in evidence]
    return values | {"done": "y", "note": _note(row)}, errors, warnings


def convert(path: Path, *, include_hold: bool) -> dict:
    """검수표 한 개 → 기준, 완료 행의 라벨 칸, 행별 보고."""
    import openpyxl  # dev 의존성. 서버 이미지에는 넣지 않는다

    wb = openpyxl.load_workbook(path, data_only=True)
    criteria = read_criteria(wb)
    out: dict = {"file": path.name, "criteria": criteria, "labels": {"tagging": {}, "build": {}}}
    report: dict = {"done": {}, "held": [], "unfinished": [], "errors": [], "warnings": []}
    sources = [("tagging", "tagging"), ("build", "build")] + ([("hold", "tagging")] if include_hold else [])
    for sheet, kind in sources:
        convert_row = tagging_values if kind == "tagging" else build_values
        done = 0
        for row in read_review(wb[SHEETS[sheet]]):
            cid, status = _text(row["문항 ID"]), _text(row.get("검수 상태"))
            if status == "보류":
                report["held"].append({"caseId": cid, "memo": _text(row.get("메모"))})
                continue
            if status != "완료":
                report["unfinished"].append(cid)
                continue
            values, errors, warnings = convert_row(row, criteria["version"])
            if warnings:
                report["warnings"].append({"caseId": cid, "warnings": warnings})
            if errors:
                report["errors"].append({"caseId": cid, "errors": errors})
                continue
            out["labels"][kind][cid] = values
            done += 1
        report["done"][SHEETS[sheet]] = done
    if not include_hold:
        report["holdSheetSkipped"] = len(read_review(wb[SHEETS["hold"]]))
    out["report"] = report
    return out


def write_label_sheet(path: Path, kind: str, items: list[dict], labels: dict[str, dict]) -> None:
    """gold_drafts.py export 와 같은 머리글·행 순서로 쓰고 라벨 칸만 채운다."""
    header, rows = (tagging_sheet_rows if kind == "tagging" else build_sheet_rows)(items)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=header)
        writer.writeheader()
        for row in rows:
            record = dict(zip(header, row, strict=True))
            writer.writerow(record | labels.get(record["case_id"], {}))


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("workbooks", nargs="+", type=Path, help="검수자별 엑셀 (보통 A·B 두 개)")
    parser.add_argument("--include-hold", action="store_true", help="보류검수 3문항도 태깅 시트에 넣기 (합의 후)")
    parser.add_argument("--work-dir", type=Path, default=EXPORT_WORK_DIR, help="gold_drafts.py export 의 작업 폴더")
    parser.add_argument("--out-dir", type=Path, default=None, help="기본 <work-dir>/labels")
    args = parser.parse_args(argv)

    items = {}
    for kind in ("tagging", "build"):
        path = args.work_dir / kind / "items.jsonl"
        if not path.exists():
            print(
                f"{path} 가 없습니다. 먼저: uv run python scripts/gold_drafts.py export --kind {kind}", file=sys.stderr
            )
            return 2
        items[kind] = _rows(path)

    results = [convert(p, include_hold=args.include_hold) for p in args.workbooks]
    problems = [f"{r['file']}: {e}" for r in results for e in criteria_errors(r["criteria"])]
    first = results[0]["criteria"]
    for r in results[1:]:
        c = r["criteria"]
        if (c["version"], c["choices"]) != (first["version"], first["choices"]):
            problems.append(f"{r['file']}: 기준 선택·합의 버전이 {results[0]['file']} 와 다름")
    reviewers = [r["criteria"]["reviewer"] for r in results]
    if len(set(reviewers)) != len(reviewers):
        problems.append(f"검수자 ID 가 겹침: {reviewers}")
    for r in results:
        for kind in ("tagging", "build"):
            known = {i["caseId"] for i in items[kind]}
            unknown = sorted(set(r["labels"][kind]) - known)
            if unknown:
                problems.append(f"{r['file']}: export 목록에 없는 {kind} 문항 {unknown[:5]}")
    if problems:
        print("\n".join(problems), file=sys.stderr)
        return 2

    out_dir = args.out_dir or args.work_dir / "labels"
    summary = {"criteria": first, "out": str(out_dir), "reviewers": {}}
    for r in results:
        reviewer = r["criteria"]["reviewer"]
        for kind in ("tagging", "build"):
            write_label_sheet(out_dir / f"{kind}_{reviewer}.csv", kind, items[kind], r["labels"][kind])
        summary["reviewers"][reviewer] = r["report"]
    (out_dir / "convert_report.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")

    for reviewer, report in summary["reviewers"].items():
        brief = {k: (len(v) if isinstance(v, list) else v) for k, v in report.items()}
        print(reviewer, json.dumps(brief, ensure_ascii=False))
        for e in report["errors"]:
            print(f"  오류 {e['caseId']}: {'; '.join(e['errors'])}", file=sys.stderr)
    print(f"보고서: {out_dir / 'convert_report.json'}")
    return 1 if any(r["report"]["errors"] for r in results) else 0


if __name__ == "__main__":
    raise SystemExit(main())
